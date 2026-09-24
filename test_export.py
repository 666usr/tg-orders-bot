import os
import tempfile
import unittest

from openpyxl import load_workbook

import export
import storage


class TestExport(unittest.TestCase):
    def setUp(self):
        handle, self.db = tempfile.mkstemp(suffix=".db")
        os.close(handle)
        storage.init_db(self.db)

        handle, self.xlsx = tempfile.mkstemp(suffix=".xlsx")
        os.close(handle)

    def tearDown(self):
        os.unlink(self.db)
        os.unlink(self.xlsx)

    def test_export_creates_file_with_headers(self):
        storage.add_order(1, "a", "Анна", "@anna", "Нужен бот", db_path=self.db)
        export.export_orders_to_xlsx(self.xlsx, db_path=self.db)

        sheet = load_workbook(self.xlsx).active
        self.assertEqual(sheet.cell(row=1, column=1).value, "№")
        self.assertEqual(sheet.max_row, 2)  # заголовок + одна заявка

    def test_export_includes_order_data(self):
        storage.add_order(1, "a", "Борис", "@boris", "Сделать парсер", db_path=self.db)
        export.export_orders_to_xlsx(self.xlsx, db_path=self.db)

        sheet = load_workbook(self.xlsx).active
        self.assertEqual(sheet.cell(row=2, column=3).value, "Борис")
        self.assertEqual(sheet.cell(row=2, column=6).value, "новая")


if __name__ == "__main__":
    unittest.main()
